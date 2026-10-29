"""엑셀·CSV 한 장의 상품 목록을 검수해 오류를 행·열별로 알려 주고, 통과한 상품만 정리된 CSV로 내보낸다. **업로드는 하지 않는다.**
검수 기준(제목 50자, 금지어, 가격 단위 등)은 예시이며 실제 스마트스토어 규칙과 다를 수 있다."""
import csv, itertools, json, os, re

HEADERS = {"상품명": "title", "판매가": "price", "재고": "stock", "카테고리": "category", "옵션": "options", "대표이미지": "image", "상세설명": "desc", "배송비": "shipping", "브랜드": "brand"}
REQUIRED = ["title", "price", "stock", "category", "image", "desc"]
TITLE_MAX, DESC_MIN, MAX_COMBOS = 50, 20, 100
BAD_WORDS = ["최저가", "1위", "무조건", "100%", "완치", "기적"]
BAD_CHARS = re.compile(r"[★☆♥♡◆◇※■□●○▶◀]")


class ListError(ValueError):
    pass


def read(path):
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".xlsx":
            from openpyxl import load_workbook
            ws = load_workbook(path, read_only=True, data_only=True).active; rows = [[("" if c is None else str(c).strip()) for c in r] for r in ws.iter_rows(values_only=True)]
        else:
            rows = list(csv.reader(open(path, encoding="utf-8-sig", newline="")))
    except (OSError, KeyError, ValueError) as e: raise ListError(f"파일을 읽지 못했습니다: {e}")
    if not rows: raise ListError("빈 파일이에요")
    head = [h.strip() for h in rows[0]]; unknown = [h for h in head if h and h not in HEADERS]; missing = [k for k, v in HEADERS.items() if v in REQUIRED and k not in head]
    if missing: raise ListError(f"필수 열이 없어요: {missing} (있는 열: {head})")
    out = []
    for i, r in enumerate(rows[1:], 2):
        if not any(c.strip() for c in r): continue
        out.append({"row": i, **{HEADERS[h]: (r[j].strip() if j < len(r) else "") for j, h in enumerate(head) if h in HEADERS}})
    return out, unknown


def parse_options(s):
    """'색상:빨강,파랑;사이즈:S,M' → {'색상':['빨강','파랑'],'사이즈':['S','M']}"""
    out = {}
    for part in [p for p in s.split(";") if p.strip()]:
        if ":" not in part: raise ListError(f"옵션 형식 오류 '{part}' (이름:값,값 형태)")
        k, v = part.split(":", 1); vals = [x.strip() for x in v.split(",") if x.strip()]
        if not k.strip() or not vals: raise ListError(f"옵션 형식 오류 '{part}'")
        if len(set(vals)) != len(vals): raise ListError(f"옵션 '{k}'에 같은 값이 겹쳐요")
        out[k.strip()] = vals
    return out


def validate(row, categories, base_dir=".", seen=None):
    errs = []; add = lambda col, msg: errs.append({"row": row["row"], "col": col, "msg": msg})
    for k in REQUIRED:
        if not row.get(k): add(k, "비어 있어요")
    t = row.get("title", "")
    if t:
        if len(t) > TITLE_MAX: add("title", f"상품명이 {len(t)}자예요({TITLE_MAX}자 이하)")
        if BAD_CHARS.search(t): add("title", "특수문자(★♥ 등)는 쓸 수 없어요")
        for w in BAD_WORDS:
            if w in t: add("title", f"'{w}' 표현은 쓸 수 없어요")
        words = t.split()
        if len(words) != len(set(words)): add("title", "같은 단어가 반복돼요(검색 어뷰징으로 보일 수 있어요)")
        if seen is not None:
            if t in seen: add("title", f"{seen[t]}행과 상품명이 같아요")
            else: seen[t] = row["row"]
    if row.get("price"):
        if not re.fullmatch(r"\d[\d,]*", row["price"]): add("price", "가격은 숫자만 적어 주세요")
        else:
            v = int(row["price"].replace(",", ""))
            if v < 100: add("price", "가격이 100원 미만이에요")
            elif v % 10: add("price", "가격은 10원 단위여야 해요")
    if row.get("stock") and not re.fullmatch(r"\d+", row["stock"]): add("stock", "재고는 0 이상의 정수여야 해요")
    if row.get("category") and row["category"] not in categories: add("category", f"모르는 카테고리 코드예요: {row['category']}")
    if row.get("image"):
        im = row["image"]
        if im.startswith("http://"): add("image", "http 주소는 안 돼요(https만)")
        elif not im.startswith("https://") and not os.path.exists(os.path.join(base_dir, im)): add("image", f"이미지 파일이 없어요: {im}")
    if row.get("desc") and len(row["desc"]) < DESC_MIN: add("desc", f"상세설명이 너무 짧아요({DESC_MIN}자 이상)")
    if row.get("shipping") and not re.fullmatch(r"\d+", row["shipping"]): add("shipping", "배송비는 숫자(원)만 적어 주세요")
    if row.get("options"):
        try:
            o = parse_options(row["options"]); n = len(list(itertools.product(*o.values())))
            if n > MAX_COMBOS: add("options", f"옵션 조합이 {n}개예요({MAX_COMBOS}개 이하)")
        except ListError as e: add("options", str(e))
    return errs


def run(path, categories, base_dir=None):
    rows, unknown = read(path); base_dir = base_dir or os.path.dirname(os.path.abspath(path)); seen = {}; good, errors = [], []
    for r in rows:
        e = validate(r, categories, base_dir, seen); (errors.extend(e) if e else good.append(r))
    return {"total": len(rows), "good": good, "errors": errors, "bad_rows": len({e["row"] for e in errors}), "unknown_columns": unknown}


def export(good, out):
    cols = ["title", "price", "stock", "category", "options", "image", "desc", "shipping", "brand"]
    with open(out, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow([k for c in cols for k, v in HEADERS.items() if v == c])
        for r in good: w.writerow([("'" + r.get(c, "") if r.get(c, "") and r.get(c, "")[0] in "=+-@" else r.get(c, "")) for c in cols])      # 엑셀 수식 방지(빈 칸은 그대로)
