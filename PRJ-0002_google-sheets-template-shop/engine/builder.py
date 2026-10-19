"""구글 시트에서 열 수 있는 수식 완성형 .xlsx 템플릿 생성기 (openpyxl)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.utils import get_column_letter
from .presets import PRESETS, check

THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


class Style:
    def __init__(self, theme):
        t = PRESETS["themes"][theme]
        self.head_fill = PatternFill("solid", fgColor=t["main"])
        self.sub_fill = PatternFill("solid", fgColor=t["sub"])
        self.head_font = Font(bold=True, color="FFFFFF")
        self.main = t["main"]


def _header(ws, row, names, st, widths=None):
    for i, n in enumerate(names, start=1):
        c = ws.cell(row=row, column=i, value=n)
        c.fill, c.font, c.border = st.head_fill, st.head_font, BORDER
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 16)


def _guide(wb, st, title, lines):
    ws = wb.create_sheet("사용설명서", 0)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16, color=st.main)
    for i, l in enumerate(lines, start=3):
        ws.cell(row=i, column=1, value=l)
    ws.column_dimensions["A"].width = 90
    return ws


# ---------------- 가계부 ----------------
def build_budget(wb, job, st):
    J = PRESETS["jobs"][job]
    cats = J["income"] + J["expense"]
    s = wb.create_sheet("설정")
    _header(s, 1, ["구분", "분류"], st)
    for i, c in enumerate(J["income"], start=2):
        s.cell(i, 1, "수입")
        s.cell(i, 2, c)
    for j, c in enumerate(J["expense"], start=2 + len(J["income"])):
        s.cell(j, 1, "지출")
        s.cell(j, 2, c)
    cat_end = 1 + len(cats)

    t = wb.create_sheet("거래내역")
    _header(t, 1, ["날짜", "구분", "분류", "내용", "금액", "월"], st, {1: 12, 2: 8, 3: 16, 4: 28, 5: 14, 6: 8})
    sample = [("2027-01-05", "수입", J["income"][0], "1월 수입", 3000000),
              ("2027-01-07", "지출", J["expense"][0], "장보기", 85000),
              ("2027-01-15", "지출", J["expense"][1], "교통비", 60000),
              ("2027-02-03", "수입", J["income"][0], "2월 수입", 2800000),
              ("2027-02-10", "지출", J["expense"][0], "외식", 120000)]
    from datetime import date
    for r, (d, k, c, m, a) in enumerate(sample, start=2):
        t.cell(r, 1, date.fromisoformat(d)).number_format = "yyyy-mm-dd"
        t.cell(r, 2, k)
        t.cell(r, 3, c)
        t.cell(r, 4, m)
        t.cell(r, 5, a).number_format = "#,##0"
    LAST = 500
    for r in range(2, LAST + 1):
        t.cell(r, 6, f'=IF(A{r}="","",MONTH(A{r}))')
        t.cell(r, 5).number_format = "#,##0"
    dv1 = DataValidation(type="list", formula1='"수입,지출"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1=f"=설정!$B$2:$B${cat_end}", allow_blank=True)
    t.add_data_validation(dv1)
    t.add_data_validation(dv2)
    dv1.add(f"B2:B{LAST}")
    dv2.add(f"C2:C{LAST}")
    t.freeze_panes = "A2"

    m = wb.create_sheet("월별요약")
    _header(m, 1, ["월", "수입", "지출", "잔액", "저축률"], st)
    for mo in range(1, 13):
        r = mo + 1
        m.cell(r, 1, mo)
        m.cell(r, 2, f'=SUMIFS(거래내역!$E$2:$E${LAST},거래내역!$F$2:$F${LAST},A{r},거래내역!$B$2:$B${LAST},"수입")')
        m.cell(r, 3, f'=SUMIFS(거래내역!$E$2:$E${LAST},거래내역!$F$2:$F${LAST},A{r},거래내역!$B$2:$B${LAST},"지출")')
        m.cell(r, 4, f"=B{r}-C{r}")
        m.cell(r, 5, f'=IF(B{r}=0,0,D{r}/B{r})').number_format = "0.0%"
        for c in (2, 3, 4):
            m.cell(r, c).number_format = "#,##0"
    m.cell(14, 1, "합계").font = Font(bold=True)
    for c, col in ((2, "B"), (3, "C"), (4, "D")):
        m.cell(14, c, f"=SUM({col}2:{col}13)").number_format = "#,##0"
    m.cell(14, 5, '=IF(B14=0,0,D14/B14)').number_format = "0.0%"
    m.conditional_formatting.add("D2:D13", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))
    ch = BarChart()
    ch.title = "월별 수입·지출"
    ch.height, ch.width = 8, 16
    ch.add_data(Reference(m, min_col=2, max_col=3, min_row=1, max_row=13), titles_from_data=True)
    ch.set_categories(Reference(m, min_col=1, min_row=2, max_row=13))
    m.add_chart(ch, "G2")

    c = wb.create_sheet("분류별지출")
    _header(c, 1, ["분류", "지출 합계", "비율"], st)
    n = len(J["expense"])
    for i, name in enumerate(J["expense"], start=2):
        c.cell(i, 1, name)
        c.cell(i, 2, f'=SUMIFS(거래내역!$E$2:$E${LAST},거래내역!$C$2:$C${LAST},A{i},거래내역!$B$2:$B${LAST},"지출")').number_format = "#,##0"
        c.cell(i, 3, f"=IF(SUM($B$2:$B${n+1})=0,0,B{i}/SUM($B$2:$B${n+1}))").number_format = "0.0%"
    pie = PieChart()
    pie.title = "분류별 지출 비율"
    pie.height, pie.width = 8, 10
    pie.add_data(Reference(c, min_col=2, min_row=1, max_row=n + 1), titles_from_data=True)
    pie.set_categories(Reference(c, min_col=1, min_row=2, max_row=n + 1))
    c.add_chart(pie, "E2")
    return ["거래내역 시트에 날짜·구분·분류·내용·금액을 입력합니다.",
            "구분과 분류는 목록에서 선택합니다. 분류는 설정 시트에서 바꿀 수 있습니다.",
            "월별요약과 분류별지출은 자동으로 계산됩니다. 수식 칸은 지우지 마세요.",
            f"거래내역은 {LAST-1}행까지 계산됩니다."]


# ---------------- 프로젝트 관리 ----------------
def build_project(wb, job, st):
    J = PRESETS["jobs"][job]
    stages = J["project_stages"]
    s = wb.create_sheet("설정")
    _header(s, 1, ["단계"], st)
    for i, x in enumerate(stages, start=2):
        s.cell(i, 1, x)
    t = wb.create_sheet("작업목록")
    _header(t, 1, ["번호", "작업", "담당", "단계", "시작일", "마감일", "진행률", "남은일수", "상태"], st,
            {1: 6, 2: 30, 3: 12, 4: 12, 5: 12, 6: 12, 7: 10, 8: 10, 9: 12})
    from datetime import date, timedelta
    base = date(2027, 1, 4)
    sample = [("기획서 작성", "나", stages[0], 0, 5, 1.0),
              ("1차 작업", "나", stages[2], 5, 14, 0.5),
              ("검토 요청", "팀", stages[3], 14, 18, 0.0),
              ("마무리", "나", stages[-1], 18, 21, 0.0)]
    LAST = 200
    for r, (n, who, stg, a, b, p) in enumerate(sample, start=2):
        t.cell(r, 2, n)
        t.cell(r, 3, who)
        t.cell(r, 4, stg)
        t.cell(r, 5, base + timedelta(days=a))
        t.cell(r, 6, base + timedelta(days=b))
        t.cell(r, 7, p)
    for r in range(2, LAST + 1):
        t.cell(r, 1, f'=IF(B{r}="","",ROW()-1)')
        t.cell(r, 5).number_format = t.cell(r, 6).number_format = "yyyy-mm-dd"
        t.cell(r, 7).number_format = "0%"
        t.cell(r, 8, f'=IF(OR(F{r}="",G{r}>=1),"",F{r}-TODAY())')
        t.cell(r, 9, f'=IF(B{r}="","",IF(G{r}>=1,"완료",IF(AND(F{r}<>"",F{r}<TODAY()),"지연",IF(G{r}>0,"진행중","대기"))))')
    dv = DataValidation(type="list", formula1=f"=설정!$A$2:$A${1+len(stages)}", allow_blank=True)
    t.add_data_validation(dv)
    dv.add(f"D2:D{LAST}")
    dvp = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                         error="진행률은 0%~100% 사이로 입력하세요")
    t.add_data_validation(dvp)
    dvp.add(f"G2:G{LAST}")
    t.conditional_formatting.add(f"I2:I{LAST}", FormulaRule(formula=['I2="지연"'], fill=PatternFill("solid", bgColor="F8CBAD"), font=Font(color="C00000", bold=True)))
    t.conditional_formatting.add(f"I2:I{LAST}", FormulaRule(formula=['I2="완료"'], fill=PatternFill("solid", bgColor="C6E0B4")))
    t.freeze_panes = "A2"
    d = wb.create_sheet("대시보드")
    _header(d, 1, ["항목", "값"], st, {1: 20, 2: 14})
    rows = [("전체 작업", f'=COUNTA(작업목록!$B$2:$B${LAST})'),
            ("완료", f'=COUNTIF(작업목록!$I$2:$I${LAST},"완료")'),
            ("진행중", f'=COUNTIF(작업목록!$I$2:$I${LAST},"진행중")'),
            ("대기", f'=COUNTIF(작업목록!$I$2:$I${LAST},"대기")'),
            ("지연", f'=COUNTIF(작업목록!$I$2:$I${LAST},"지연")'),
            ("전체 진행률", f'=IF(B2=0,0,SUM(작업목록!$G$2:$G${LAST})/B2)')]
    for i, (k, f) in enumerate(rows, start=2):
        d.cell(i, 1, k)
        d.cell(i, 2, f)
    d["B7"].number_format = "0.0%"
    ch = BarChart()
    ch.title = "상태별 작업 수"
    ch.height, ch.width = 7, 12
    ch.add_data(Reference(d, min_col=2, min_row=3, max_row=6))
    ch.set_categories(Reference(d, min_col=1, min_row=3, max_row=6))
    ch.legend = None
    d.add_chart(ch, "D2")
    return ["작업목록 시트에 작업·담당·단계·시작일·마감일·진행률을 입력합니다.",
            "남은일수와 상태는 자동입니다. 상태는 진행률과 마감일로 계산됩니다(지연/완료/진행중/대기).",
            "대시보드에서 전체 현황을 봅니다."]


# ---------------- 재고 관리 ----------------
def build_inventory(wb, job, st):
    J = PRESETS["jobs"][job]
    s = wb.create_sheet("설정")
    _header(s, 1, ["단위"], st)
    for i, x in enumerate(J["inventory_units"], start=2):
        s.cell(i, 1, x)
    p = wb.create_sheet("상품목록")
    _header(p, 1, ["상품코드", "상품명", "단위", "단가", "안전재고"], st, {1: 12, 2: 26, 3: 8, 4: 12, 5: 10})
    prods = [("A001", "상품 A", 1500, 20), ("A002", "상품 B", 3200, 10), ("A003", "상품 C", 800, 50)]
    for r, (c, n, pr, sf) in enumerate(prods, start=2):
        p.cell(r, 1, c)
        p.cell(r, 2, n)
        p.cell(r, 3, J["inventory_units"][0])
        p.cell(r, 4, pr).number_format = "#,##0"
        p.cell(r, 5, sf)
    dv = DataValidation(type="list", formula1=f"=설정!$A$2:$A${1+len(J['inventory_units'])}", allow_blank=True)
    p.add_data_validation(dv)
    dv.add("C2:C200")
    io = wb.create_sheet("입출고")
    _header(io, 1, ["날짜", "상품코드", "구분", "수량"], st, {1: 12, 2: 12, 3: 8, 4: 10})
    from datetime import date
    rows = [("2027-01-04", "A001", "입고", 100), ("2027-01-04", "A002", "입고", 40), ("2027-01-04", "A003", "입고", 120),
            ("2027-01-09", "A001", "출고", 85), ("2027-01-12", "A002", "출고", 33), ("2027-01-15", "A003", "출고", 60)]
    for r, (d, c, k, q) in enumerate(rows, start=2):
        io.cell(r, 1, date.fromisoformat(d)).number_format = "yyyy-mm-dd"
        io.cell(r, 2, c)
        io.cell(r, 3, k)
        io.cell(r, 4, q)
    LAST = 1000
    dv1 = DataValidation(type="list", formula1='"입고,출고"', allow_blank=True)
    io.add_data_validation(dv1)
    dv1.add(f"C2:C{LAST}")
    dv2 = DataValidation(type="list", formula1="=상품목록!$A$2:$A$200", allow_blank=True)
    io.add_data_validation(dv2)
    dv2.add(f"B2:B{LAST}")
    io.freeze_panes = "A2"
    r = wb.create_sheet("재고현황")
    _header(r, 1, ["상품코드", "상품명", "입고합계", "출고합계", "현재재고", "재고금액", "상태"], st, {2: 24})
    for i in range(2, 202):
        r.cell(i, 1, f'=IF(상품목록!A{i}="","",상품목록!A{i})')
        r.cell(i, 2, f'=IF(A{i}="","",상품목록!B{i})')
        r.cell(i, 3, f'=IF(A{i}="","",SUMIFS(입출고!$D$2:$D${LAST},입출고!$B$2:$B${LAST},A{i},입출고!$C$2:$C${LAST},"입고"))')
        r.cell(i, 4, f'=IF(A{i}="","",SUMIFS(입출고!$D$2:$D${LAST},입출고!$B$2:$B${LAST},A{i},입출고!$C$2:$C${LAST},"출고"))')
        r.cell(i, 5, f'=IF(A{i}="","",C{i}-D{i})')
        r.cell(i, 6, f'=IF(A{i}="","",E{i}*상품목록!D{i})').number_format = "#,##0"
        r.cell(i, 7, f'=IF(A{i}="","",IF(E{i}<0,"오류",IF(E{i}<=상품목록!E{i},"발주 필요","정상")))')
    r.conditional_formatting.add("G2:G201", FormulaRule(formula=['G2="발주 필요"'], fill=PatternFill("solid", bgColor="FFE699"), font=Font(bold=True)))
    r.conditional_formatting.add("G2:G201", FormulaRule(formula=['G2="오류"'], fill=PatternFill("solid", bgColor="F8CBAD"), font=Font(color="C00000", bold=True)))
    r.freeze_panes = "A2"
    return ["상품목록 시트에 상품코드·상품명·단가·안전재고를 등록합니다.",
            "입출고 시트에 날짜·상품코드·구분(입고/출고)·수량을 입력합니다.",
            "재고현황은 자동으로 계산됩니다. 현재재고가 안전재고 이하이면 '발주 필요'가 표시됩니다.",
            "상품은 200개까지, 입출고는 999행까지 계산됩니다."]


TITLES = {"budget": "가계부", "project": "프로젝트 관리", "inventory": "재고 관리"}
BUILDERS = {"budget": build_budget, "project": build_project, "inventory": build_inventory}


def build_template(kind, job="freelancer", theme="navy", out_dir="output"):
    check(kind, job, theme)
    st = Style(theme)
    wb = Workbook()
    wb.remove(wb.active)
    lines = BUILDERS[kind](wb, job, st)
    label = PRESETS["jobs"][job]["label"]["ko"]
    _guide(wb, st, f"{TITLES[kind]} 템플릿 ({label}용)",
           lines + ["", "※ 흰색 칸에 입력하세요. 수식이 들어 있는 칸은 수정하지 마세요.",
                    "※ 구글 시트에서는 [파일 > 사본 만들기]로 사용하세요."])
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = (ws.title != "사용설명서")
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"{kind}_{job}_{theme}.xlsx")
    wb.save(path)
    return path
