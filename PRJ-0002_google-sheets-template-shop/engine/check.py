"""수식 검사기: (1) 정적 검사 (2) LibreOffice로 다시 계산해 오류값 찾기."""
import os
import re
import shutil
import subprocess
import tempfile
from openpyxl import load_workbook

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "Err:")
SHEET_REF = re.compile(r"(?:'([^']+)'|([A-Za-z0-9_가-힣]+))!")


def static_check(path):
    """존재하지 않는 시트 참조, #REF! 문자열을 찾는다."""
    wb = load_workbook(path)
    names = set(wb.sheetnames)
    problems = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    if "#REF!" in v:
                        problems.append((ws.title, c.coordinate, "#REF! 포함"))
                    for m in SHEET_REF.finditer(v):
                        nm = m.group(1) or m.group(2)
                        if nm not in names:
                            problems.append((ws.title, c.coordinate, f"없는 시트: {nm}"))
    return problems


def recalc(path, out_dir=None):
    """LibreOffice로 다시 계산한 복사본 경로를 돌려준다. 없으면 None."""
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    out_dir = out_dir or tempfile.mkdtemp(prefix="recalc_")
    prof = tempfile.mkdtemp(prefix="lo_profile_")
    subprocess.run([exe, f"-env:UserInstallation=file://{prof}", "--headless", "--calc",
                    "--convert-to", "xlsx", "--outdir", out_dir, path],
                   capture_output=True, timeout=180)
    out = os.path.join(out_dir, os.path.basename(path))
    return out if os.path.exists(out) else None


def error_values(recalculated):
    wb = load_workbook(recalculated, data_only=True)
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERRORS):
                    bad.append((ws.title, c.coordinate, c.value))
    return bad


def check_workbook(path):
    res = {"static": static_check(path), "errors": [], "recalculated": False}
    rc = recalc(path)
    if rc:
        res["recalculated"] = True
        res["errors"] = error_values(rc)
        res["recalc_path"] = rc
    res["ok"] = not res["static"] and not res["errors"] and res["recalculated"]
    return res
